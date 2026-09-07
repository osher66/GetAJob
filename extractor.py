"""
extractor.py - מודול חילוץ טקסט מקובצי קורות חיים (PDF ו-DOCX)
תומך בחילוץ ישיר בזיכרון (In-Memory BytesIO) ללא שמירה בדיסק,
כולל אימות שלמות הקובץ ובדיקת קריאות טקסט.
"""

import io
from typing import Tuple
from pypdf import PdfReader

try:
    import docx
except ImportError:
    docx = None


class DocumentExtractionError(Exception):
    """שגיאה מותאמת עבור בעיות בחילוץ קובץ מסמך"""
    pass


# תאימות לאחור
PDFExtractionError = DocumentExtractionError


def extract_text_from_pdf(file_bytes: bytes) -> Tuple[str, int]:
    """
    מחלץ טקסט מקובץ PDF ישירות מזרם בייטים בזיכרון.
    מחזיר: (טקסט מחולץ, מספר עמודים)
    """
    if not file_bytes:
        raise DocumentExtractionError("הקובץ שהועלה ריק.")

    try:
        pdf_stream = io.BytesIO(file_bytes)
        reader = PdfReader(pdf_stream)
        num_pages = len(reader.pages)

        extracted_pages = []
        for page in reader.pages:
            page_text = page.extract_text() or ""
            extracted_pages.append(page_text)

        full_text = "\n".join(extracted_pages).strip()

        if len(full_text) < 30:
            raise DocumentExtractionError(
                "הקובץ אינו מכיל שכבת טקסט קריאה (ייתכן שמדובר בקובץ סרוק או תמונה). "
                "אנא העלה קובץ PDF עם טקסט דיגיטלי."
            )

        return full_text, num_pages

    except DocumentExtractionError:
        raise
    except Exception as e:
        raise DocumentExtractionError(f"שגיאה בעיבוד קובץ ה-PDF: {str(e)}")


def extract_text_from_docx(file_bytes: bytes) -> Tuple[str, int]:
    """
    מחלץ טקסט מקובץ Word (.docx) ישירות מזרם בייטים בזיכרון.
    מחזיר: (טקסט מחולץ, מספר פסקאות)
    """
    if not file_bytes:
        raise DocumentExtractionError("הקובץ שהועלה ריק.")

    if docx is None:
        raise DocumentExtractionError(
            "ספריית python-docx אינה מותקנת בשרת. אנא העלה קובץ בפורמט PDF."
        )

    try:
        doc_stream = io.BytesIO(file_bytes)
        doc = docx.Document(doc_stream)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

        # חילוץ טקסט גם מתוך טבלאות אם קיימות
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    if cell_text and cell_text not in paragraphs:
                        paragraphs.append(cell_text)

        full_text = "\n".join(paragraphs).strip()

        if len(full_text) < 30:
            raise DocumentExtractionError(
                "קובץ ה-Word שהועלה ריק או מכיל כמות טקסט מזערית בלבד."
            )

        return full_text, len(paragraphs)

    except DocumentExtractionError:
        raise
    except Exception as e:
        raise DocumentExtractionError(f"שגיאה בעיבוד קובץ ה-DOCX: {str(e)}")


def extract_text_from_file(file_bytes: bytes, filename: str) -> Tuple[str, int]:
    """
    מזהה את סוג הקובץ לפי הסיומת ומחלץ את הטקסט בהתאם.
    """
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "pdf":
        return extract_text_from_pdf(file_bytes)
    elif ext in ("docx", "doc"):
        if ext == "doc":
            raise DocumentExtractionError("קובצי .doc ישנים אינם נתמכים ישירות. אנא שמור כקובץ .docx או .pdf.")
        return extract_text_from_docx(file_bytes)
    else:
        raise DocumentExtractionError(f"פורמט קובץ לא נתמך (.{ext}). אנא העלה קובץ PDF או DOCX.")
